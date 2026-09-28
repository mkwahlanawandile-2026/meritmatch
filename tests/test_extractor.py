import pytest
from docx import Document

import config
from src.extractor import UnsupportedFileError, clean_text, extract_text


def test_clean_text_normalizes_whitespace_and_bullets():
    raw = "Skills:\u2022  Python\n\n\n\nSQL"
    assert clean_text(raw) == "Skills:- Python\n\nSQL"


def test_extract_txt(tmp_path):
    f = tmp_path / "resume.txt"
    f.write_text("Jane Doe\nPython developer", encoding="utf-8")
    assert "Python developer" in extract_text(f)


def test_extract_docx_including_tables(tmp_path):
    doc = Document()
    doc.add_paragraph("John Smith")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Skills: Machine Learning"
    f = tmp_path / "resume.docx"
    doc.save(f)
    text = extract_text(f)
    assert "John Smith" in text
    assert "Machine Learning" in text


def test_unsupported_type_rejected(tmp_path):
    f = tmp_path / "resume.png"
    f.write_bytes(b"not a resume")
    with pytest.raises(UnsupportedFileError):
        extract_text(f)


def test_oversized_file_rejected(tmp_path):
    f = tmp_path / "big.txt"
    f.write_bytes(b"a" * (config.MAX_FILE_SIZE_MB * 1024 * 1024 + 1))
    with pytest.raises(UnsupportedFileError):
        extract_text(f)
