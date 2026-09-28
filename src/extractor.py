"""Extract clean text from resume files (PDF, DOCX, TXT)."""
import io
import re
from pathlib import Path

import pdfplumber
from docx import Document

import config


class UnsupportedFileError(ValueError):
    """Raised when a file type is not supported or a file is too large."""


def _read_bytes(source):
    """Return (filename, bytes) from a file path or an uploaded file object."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        return path.name, path.read_bytes()
    # Streamlit UploadedFile or any file-like object with .name
    data = source.getvalue() if hasattr(source, "getvalue") else source.read()
    return source.name, data


def _pdf_to_text(data: bytes) -> str:
    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _docx_to_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # Include text inside tables (many resumes use table layouts)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def clean_text(text: str) -> str:
    """Normalize whitespace and remove junk characters."""
    text = text.replace("\u2022", "-").replace("\u2013", "-").replace("\u2014", "-")
    text = re.sub(r"[^\x09\x0A\x20-\x7E]", " ", text)  # drop non-printable/odd chars
    text = re.sub(r"[ \t]+", " ", text)                 # collapse spaces
    text = re.sub(r"\n{3,}", "\n\n", text)              # collapse blank lines
    return text.strip()


def extract_text(source) -> str:
    """Extract clean text from a file path or uploaded file object."""
    filename, data = _read_bytes(source)
    ext = Path(filename).suffix.lower()

    if ext not in config.ALLOWED_EXTENSIONS:
        raise UnsupportedFileError(
            f"Unsupported file type '{ext}'. Allowed: {sorted(config.ALLOWED_EXTENSIONS)}"
        )
    if len(data) > config.MAX_FILE_SIZE_MB * 1024 * 1024:
        raise UnsupportedFileError(
            f"File too large. Maximum size is {config.MAX_FILE_SIZE_MB} MB."
        )

    if ext == ".pdf":
        raw = _pdf_to_text(data)
    elif ext == ".docx":
        raw = _docx_to_text(data)
    else:
        raw = data.decode("utf-8", errors="ignore")

    return clean_text(raw)
